#!/usr/bin/env python3
"""Build a sourced, live-formula workbook and its Python calculation sidecar."""
import argparse
import hashlib
import json
import math
from pathlib import Path

from formula_model import FormulaModel, MissingInput

FIELDS = "revenue cogs gross_profit opex da ebit interest ebt taxes net_income sbc capex receivables inventory other_ca ppe other_assets cash assets payables other_cl debt other_liabilities equity retained_earnings liabilities cfo cfi cff fx diluted_shares fcff nwc delta_nwc balance_check cash_check re_check ppe_check financing_gap".split()
ROWS = {name: i + 5 for i, name in enumerate(FIELDS)}
DRIVERS = "growth cogs_pct opex_pct depreciation_rate capex_pct ar_pct inventory_pct other_ca_pct ap_pct other_cl_pct tax_rate interest_rate debt_issued debt_repaid dividends buybacks equity_issued sbc oci fx cfi_other shares_issued shares_repur".split()
CASES = ("bear", "base", "bull")
INPUTS = "wacc terminal_growth debt excess_cash minority other_claims nonoperating_assets diluted_shares market_cap cost_equity terminal_roe terminal_payout".split()
SCALES = {"units": 1, "thousand": 1000, "million": 1000000, "billion": 1000000000}


def column(number):
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def number(value, path):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{path}: 必须是有限数字或null")
    return value


def provenance(data, path):
    parts = path.split(".")
    while parts:
        note = data.get("provenance", {}).get(".".join(parts))
        if note:
            return f"{path}: {note}"
        parts.pop()
    return f"UNTRACED {path}：尚未关联证据/假设"


def raw(graph, data, sheet, address, value, path, percent=False):
    return graph.set(sheet, address, number(value, path), note=provenance(data, path), percent=percent)


def formula(graph, sheet, address, expression, percent=False):
    return graph.set(sheet, address, formula="=" + expression, percent=percent)


def at(sheet, field, col):
    return f"'{sheet}'!{col}{ROWS[field]}"


def validate_input(data):
    for key in ("company", "as_of", "currency", "money_unit", "shares_unit", "industry_type", "valuation"):
        if not data.get(key):
            raise ValueError(f"缺少 {key}")
    from datetime import date
    date.fromisoformat(data["as_of"])
    for key in ("money_unit", "shares_unit"):
        if data[key] not in SCALES:
            raise ValueError(f"{key} 必须是 {list(SCALES)}")
    method = data["valuation"].get("method")
    if method not in {"fcff", "rim", "ddm", "nav", "unavailable"}:
        raise ValueError("valuation.method 必须为 fcff/rim/ddm/nav/unavailable")
    if data["industry_type"] in {"bank", "insurance"} and method == "fcff":
        raise ValueError("银行/保险不能套普通企业FCFF；请使用rim/ddm或不可估值并补行业附表")
    if method == "fcff" and (not data.get("history") or not data.get("scenarios")):
        raise ValueError("FCFF需要历史期初和三个情景；缺少资料应选unavailable")
    years = [x["year"] for x in data.get("history", [])]
    if years != sorted(set(years)) or any(isinstance(x, bool) or not isinstance(x, int) for x in years):
        raise ValueError("history年份必须按整数升序且不重复")
    if method in {"fcff", "rim", "ddm"}:
        lengths = [len(data.get("scenarios", {}).get(case, [])) for case in CASES]
        if not lengths[0] or len(set(lengths)) != 1:
            raise ValueError("三个情景必须有相同的非零预测年数")
    if data.get("selected_case", "base") not in CASES:
        raise ValueError("selected_case无效")
    if data["industry_type"] not in {"operating", "bank", "insurance", "asset", "cyclical", "early_growth", "private"}:
        raise ValueError("industry_type不在支持的路由范围")
    for i, record in enumerate(data.get("history", [])):
        for field in ("cash", "debt", "ppe", "inventory", "assets", "liabilities", "diluted_shares", "capex", "cogs", "opex"):
            value = number(record.get(field), f"history.{i}.{field}")
            if value is not None and value < 0:
                raise ValueError(f"history.{i}.{field}: 违反该基线的正值口径，需核对分类或扩展模型")
    for case, periods in data.get("scenarios", {}).items():
        for i, period in enumerate(periods):
            for field in DRIVERS:
                value = number(period.get(field), f"scenarios.{case}.{i}.{field}")
                if value is None:
                    continue
                if field == "growth" and value <= -1:
                    raise ValueError("预测增长率须大于-100%")
                if (field.endswith("pct") or field in {"depreciation_rate", "tax_rate", "interest_rate", "debt_issued", "debt_repaid", "dividends", "buybacks", "equity_issued", "sbc", "shares_issued", "shares_repur"}) and value < 0:
                    raise ValueError(f"scenarios.{case}.{i}.{field}: 需非负；净现金流使用对应带符号字段")
                if field in {"tax_rate", "depreciation_rate"} and value > 1:
                    raise ValueError(f"{field}: 超过该简化模型的适用范围，需独立附表")
    if method == "rim" and any(data["valuation"].get(x) is None for x in ("book_equity", "terminal_roe", "terminal_payout")):
        raise ValueError("RIM需要期初权益、终值ROE和可持续分红率")


def add_actuals(graph, data):
    history = data.get("history", [])
    graph.set("Actuals", "A1", "历史原始披露；未披露保持空白，不是零")
    for field, row in ROWS.items():
        graph.set("Actuals", f"A{row}", field)
    for i, record in enumerate(history):
        col = column(i + 2)
        previous = column(i + 1)
        graph.set("Actuals", f"{col}3", record["year"])
        for field, row in ROWS.items():
            raw(graph, data, "Actuals", f"{col}{row}", record.get(field), f"history.{i}.{field}")
        def f(field, expression):
            formula(graph, "Actuals", f"{col}{ROWS[field]}", expression)
        r = lambda field: at("Actuals", field, col)
        f("balance_check", f"{r('assets')}-{r('liabilities')}-{r('equity')}")
        f("nwc", f"{r('receivables')}+{r('inventory')}+{r('other_ca')}-{r('payables')}-{r('other_cl')}")
        if i and record["year"] == history[i-1]["year"] + 1:
            f("cash_check", f"{r('cash')}-{at('Actuals','cash',previous)}-{r('cfo')}-{r('cfi')}-{r('cff')}-{r('fx')}")
        f("financing_gap", f"MAX(-{r('cash')},0)")


def add_assumptions(graph, data, fields):
    graph.set("Assumptions", "A1", "假设：1=Bear，2=Base，3=Bull；active行随唯一选择器更新")
    graph.set("Assumptions", "A2", "Case selector")
    graph.set("Assumptions", "B2", CASES.index(data.get("selected_case", "base")) + 1,
              note="控制项：选择1/2/3；历史不变")
    count = len(data["scenarios"]["base"])
    last_year = (data.get("history") or [{"year": int(data["as_of"][:4])}])[-1]["year"]
    refs = {}
    for i in range(count):
        col = column(i + 2)
        graph.set("Assumptions", f"{col}3", last_year + i + 1)
        for k, field in enumerate(fields):
            row = 5 + k * 4
            percent = field.endswith("pct") or field in {"growth", "depreciation_rate", "tax_rate", "interest_rate", "roe", "payout"}
            graph.set("Assumptions", f"A{row}", f"{field} active")
            formula(graph, "Assumptions", f"{col}{row}", f"CHOOSE($B$2,{col}{row+1},{col}{row+2},{col}{row+3})", percent)
            for j, case in enumerate(CASES):
                graph.set("Assumptions", f"A{row+j+1}", f"{field} {case}")
                ref = raw(graph, data, "Assumptions", f"{col}{row+j+1}", data["scenarios"][case][i].get(field),
                          f"scenarios.{case}.{i}.{field}", percent)
                refs[case, i, field] = ref
    return refs, count


def add_operating_forecast(graph, data, sheet, case, refs, count):
    graph.set(sheet, "A1", f"{case}完整三表；期初债务计息、无并购/递延税/自动融资；SBC按经济费用估值")
    for field, row in ROWS.items():
        graph.set(sheet, f"A{row}", field)
    last = column(len(data["history"]) + 1)
    for i in range(count):
        col = column(i + 2)
        prior = column(i + 1)
        graph.set(sheet, f"{col}3", data["history"][-1]["year"] + i + 1)
        r = lambda field: at(sheet, field, col)
        p = lambda field: at(sheet, field, prior) if i else at("Actuals", field, last)
        a = lambda field: refs[case, i, field]
        def f(field, expression):
            return formula(graph, sheet, f"{col}{ROWS[field]}", expression)
        f("revenue", f"{p('revenue')}*(1+{a('growth')})")
        f("cogs", f"{r('revenue')}*{a('cogs_pct')}")
        f("gross_profit", f"{r('revenue')}-{r('cogs')}")
        f("opex", f"{r('revenue')}*{a('opex_pct')}")
        f("da", f"{p('ppe')}*{a('depreciation_rate')}")
        f("ebit", f"{r('gross_profit')}-{r('opex')}-{r('da')}")
        f("interest", f"{p('debt')}*{a('interest_rate')}")
        f("ebt", f"{r('ebit')}-{r('interest')}")
        f("taxes", f"MAX({r('ebt')},0)*{a('tax_rate')}")
        f("net_income", f"{r('ebt')}-{r('taxes')}")
        f("sbc", a("sbc"))
        f("capex", f"{r('revenue')}*{a('capex_pct')}")
        for field, driver in (("receivables", "ar_pct"), ("inventory", "inventory_pct"), ("other_ca", "other_ca_pct"),
                              ("payables", "ap_pct"), ("other_cl", "other_cl_pct")):
            f(field, f"{r('revenue')}*{a(driver)}")
        f("nwc", f"{r('receivables')}+{r('inventory')}+{r('other_ca')}-{r('payables')}-{r('other_cl')}")
        f("delta_nwc", f"{r('nwc')}-{p('nwc')}")
        f("ppe", f"{p('ppe')}+{r('capex')}-{r('da')}")
        f("other_assets", p("other_assets"))
        f("other_liabilities", p("other_liabilities"))
        f("debt", f"{p('debt')}+{a('debt_issued')}-{a('debt_repaid')}")
        f("cfo", f"{r('net_income')}+{r('da')}+{r('sbc')}-{r('delta_nwc')}")
        f("cfi", f"-{r('capex')}+{a('cfi_other')}")
        f("cff", f"{a('debt_issued')}-{a('debt_repaid')}+{a('equity_issued')}-{a('buybacks')}-{a('dividends')}")
        f("fx", a("fx"))
        f("cash", f"{p('cash')}+{r('cfo')}+{r('cfi')}+{r('cff')}+{r('fx')}")
        f("assets", f"{r('cash')}+{r('receivables')}+{r('inventory')}+{r('other_ca')}+{r('ppe')}+{r('other_assets')}")
        f("equity", f"{p('equity')}+{r('net_income')}+{r('sbc')}+{a('equity_issued')}-{a('buybacks')}-{a('dividends')}+{a('oci')}")
        f("retained_earnings", f"{p('retained_earnings')}+{r('net_income')}-{a('dividends')}")
        f("liabilities", f"{r('payables')}+{r('other_cl')}+{r('debt')}+{r('other_liabilities')}")
        f("diluted_shares", f"{p('diluted_shares')}+{a('shares_issued')}-{a('shares_repur')}")
        f("fcff", f"{r('ebit')}-MAX({r('ebit')},0)*{a('tax_rate')}+{r('da')}-{r('capex')}-{r('delta_nwc')}")
        f("balance_check", f"{r('assets')}-{r('liabilities')}-{r('equity')}")
        f("cash_check", f"{r('cash')}-{p('cash')}-{r('cfo')}-{r('cfi')}-{r('cff')}-{r('fx')}")
        f("re_check", f"{r('retained_earnings')}-{p('retained_earnings')}-{r('net_income')}+{a('dividends')}")
        f("ppe_check", f"{r('ppe')}-{p('ppe')}-{r('capex')}+{r('da')}")
        f("financing_gap", f"MAX(-{r('cash')},0)")


def add_valuation_inputs(graph, data):
    graph.set("Valuation", "A1", f"{data['company']} | {data['valuation']['method']} | {data['as_of']}")
    graph.set("Valuation", "A2", "所选情景")
    formula(graph, "Valuation", "B2", "'Assumptions'!B2") if data["valuation"]["method"] in {"fcff", "rim", "ddm"} else graph.set("Valuation", "B2", 2)
    refs = {}
    for i, field in enumerate(INPUTS, 4):
        graph.set("Valuation", f"A{i}", field)
        refs[field] = raw(graph, data, "Valuation", f"B{i}", data["valuation"].get(field), f"valuation.{field}", field in {"wacc", "terminal_growth", "cost_equity", "terminal_roe", "terminal_payout"})
    graph.set("Valuation", "A18", "money/share scale")
    refs["scale"] = graph.set("Valuation", "B18", SCALES[data["money_unit"]] / SCALES[data["shares_unit"]], note="单位换算常数")
    return refs


def bridge(refs, ev):
    return f"({ev})-{refs['debt']}+{refs['excess_cash']}-{refs['minority']}-{refs['other_claims']}+{refs['nonoperating_assets']}"


def dcf_expression(graph, sheet, count, rate, growth):
    flows = [at(sheet, "fcff", column(i+2)) for i in range(count)]
    pv = "+".join(f"{value}/(1+{rate})^{i+1}" for i, value in enumerate(flows))
    return f"({pv})+{flows[-1]}*(1+{growth})/({rate}-{growth})/(1+{rate})^{count}"


def add_sensitivity(graph, data, refs, count):
    graph.set("Sensitivity", "A1", "Base情景 WACC × g；完整折现重算，中心格=Base价值")
    for i, shift in enumerate((-0.02, -0.01, 0, 0.01, 0.02), 4):
        formula(graph, "Sensitivity", f"A{i}", f"{refs['wacc']}+({shift})", True)
    for j, shift in enumerate((-0.01, -0.005, 0, 0.005, 0.01), 2):
        col = column(j)
        formula(graph, "Sensitivity", f"{col}3", f"{refs['terminal_growth']}+({shift})", True)
        for i in range(4, 9):
            ev = dcf_expression(graph, "Base", count, f"'Sensitivity'!A{i}", f"'Sensitivity'!{col}3")
            expr = f"IF('Sensitivity'!A{i}>'Sensitivity'!{col}3,({bridge(refs, ev)})/{refs['diluted_shares']}*{refs['scale']},\"NM\")"
            formula(graph, "Sensitivity", f"{col}{i}", expr)


def growth_trial(graph, data, refs, assumption_refs, count, row, growth_ref):
    """FCFF rebuild for a uniform growth rate, holding Base margins/reinvestment fixed."""
    last = column(len(data["history"]) + 1)
    prev_rev = at("Actuals", "revenue", last)
    prev_ppe = at("Actuals", "ppe", last)
    prev_nwc = at("Actuals", "nwc", last)
    discounted = []
    for i in range(count):
        r = row + i
        a = lambda field: assumption_refs["base", i, field]
        def f(col, expression):
            return formula(graph, "Reverse", f"{col}{r}", expression)
        revenue = f("H", f"{prev_rev}*(1+{growth_ref})")
        dep = f("I", f"{prev_ppe}*{a('depreciation_rate')}")
        capex = f("J", f"{revenue}*{a('capex_pct')}")
        ppe = f("K", f"{prev_ppe}+{capex}-{dep}")
        nwc = f("L", f"{revenue}*({a('ar_pct')}+{a('inventory_pct')}+{a('other_ca_pct')}-{a('ap_pct')}-{a('other_cl_pct')})")
        ebit = f("M", f"{revenue}*(1-{a('cogs_pct')}-{a('opex_pct')})-{dep}")
        fcff = f("N", f"{ebit}-MAX({ebit},0)*{a('tax_rate')}+{dep}-{capex}-({nwc}-{prev_nwc})")
        discounted.append(f"{fcff}/(1+{refs['wacc']})^{i+1}")
        prev_rev, prev_ppe, prev_nwc = revenue, ppe, nwc
    ev = "+".join(discounted) + f"+{fcff}*(1+{refs['terminal_growth']})/({refs['wacc']}-{refs['terminal_growth']})/(1+{refs['wacc']})^{count}"
    return bridge(refs, ev)


def add_reverse(graph, data, refs, assumption_refs, count):
    graph.set("Reverse", "A1", "统一增长率二分求解；固定Base利润/再投资/折现/终值；无括区不报根")
    low = number(data["valuation"].get("reverse_low", -0.5), "reverse_low")
    high = number(data["valuation"].get("reverse_high", 1.5), "reverse_high")
    if not -1 < low < high:
        raise ValueError("反向估值区间须满足 -1 < low < high")
    graph.set("Reverse", "B3", low, note="数值求解下界，非经营预测")
    graph.set("Reverse", "C3", high, note="数值求解上界，非经营预测")
    lo_expr = growth_trial(graph, data, refs, assumption_refs, count, 5, "'Reverse'!B3")
    hi_expr = growth_trial(graph, data, refs, assumption_refs, count, 5+count, "'Reverse'!C3")
    formula(graph, "Reverse", "E3", f"({lo_expr})-{refs['market_cap']}")
    formula(graph, "Reverse", "F3", f"({hi_expr})-{refs['market_cap']}")
    for i in range(40):
        r = i+5
        graph.set("Reverse", f"A{r}", i+1)
        if i == 0:
            formula(graph, "Reverse", f"B{r}", "B3")
            formula(graph, "Reverse", f"C{r}", "C3")
            formula(graph, "Reverse", f"G{r}", "E3")
        else:
            p = r-1
            formula(graph, "Reverse", f"B{r}", f"IF(G{p}*F{p}<=0,B{p},D{p})")
            formula(graph, "Reverse", f"C{r}", f"IF(G{p}*F{p}<=0,D{p},C{p})")
            formula(graph, "Reverse", f"G{r}", f"IF(G{p}*F{p}<=0,G{p},F{p})")
        formula(graph, "Reverse", f"D{r}", f"(B{r}+C{r})/2", True)
        expr = growth_trial(graph, data, refs, assumption_refs, count, 5+count*(i+2), f"'Reverse'!D{r}")
        formula(graph, "Reverse", f"E{r}", expr)
        formula(graph, "Reverse", f"F{r}", f"E{r}-{refs['market_cap']}")
    formula(graph, "Reverse", "A47", 'IF(E3*F3<=0,"BRACKETED","NO_BRACKET")')
    formula(graph, "Reverse", "B47", 'IF(E3*F3<=0,D44,"NO_BRACKET")', True)
    formula(graph, "Reverse", "C47", 'IF(E3*F3<=0,F44,"NO_BRACKET")')


def add_equity_model(graph, data, refs, method):
    fields = ("roe", "payout") if method == "rim" else ("dividend",)
    assumptions, count = add_assumptions(graph, data, fields)
    results = {}
    for case in CASES:
        sheet = case.title()
        graph.set(sheet, "A1", "权益模型基本构件；另补行业经营/资本附表")
        for col, label in (("A", "预测年度"), ("B", "净利润" if method == "rim" else "总股利"), ("C", "总股利"), ("D", "剩余收益"), ("E", "期末权益")):
            graph.set(sheet, f"{col}4", label)
        if method == "rim":
            graph.set(sheet, "A3", "期初账面权益")
            prior = raw(graph, data, sheet, "B3", data["valuation"].get("book_equity"), "valuation.book_equity")
        flows = []
        for i in range(count):
            r = i+5
            graph.set(sheet, f"A{r}", (data.get("history") or [{"year": int(data["as_of"][:4])}])[-1]["year"]+i+1)
            if method == "rim":
                earning = formula(graph, sheet, f"B{r}", f"{prior}*{assumptions[case,i,'roe']}")
                dividend = formula(graph, sheet, f"C{r}", f"{earning}*{assumptions[case,i,'payout']}")
                flow = formula(graph, sheet, f"D{r}", f"{earning}-{prior}*{refs['cost_equity']}")
                prior = formula(graph, sheet, f"E{r}", f"{prior}+{earning}-{dividend}")
            else:
                flow = formula(graph, sheet, f"B{r}", assumptions[case,i,"dividend"])
            flows.append(f"{flow}/(1+{refs['cost_equity']})^{i+1}")
        pv = "+".join(flows)
        if method == "rim":
            terminal = f"{prior}*({refs['terminal_roe']}-{refs['cost_equity']})/({refs['cost_equity']}-{refs['terminal_growth']})"
            # Terminal ROE is earned on closing book equity; retained earnings must support g.
            equity = f"'{sheet}'!B3+({pv})+({terminal})/(1+{refs['cost_equity']})^{count}"
        else:
            equity = f"({pv})+{flow}*(1+{refs['terminal_growth']})/({refs['cost_equity']}-{refs['terminal_growth']})/(1+{refs['cost_equity']})^{count}"
        formula(graph, sheet, "B2", equity)
        results[case] = f"'{sheet}'!B2"
    return results


def make_graph(data):
    validate_input(data)
    graph = FormulaModel()
    add_actuals(graph, data)
    refs = add_valuation_inputs(graph, data)
    method = data["valuation"]["method"]
    case_values = {}
    if method == "fcff":
        assumptions, count = add_assumptions(graph, data, DRIVERS)
        for case in CASES:
            add_operating_forecast(graph, data, case.title(), case, assumptions, count)
            ev = dcf_expression(graph, case.title(), count, refs["wacc"], refs["terminal_growth"])
            formula(graph, case.title(), "B2", bridge(refs, ev))
            case_values[case] = f"'{case.title()}'!B2"
        add_sensitivity(graph, data, refs, count)
        if data["valuation"].get("market_cap") is not None:
            add_reverse(graph, data, refs, assumptions, count)
    elif method in {"rim", "ddm"}:
        case_values = add_equity_model(graph, data, refs, method)
    elif method == "nav":
        graph.set("NAV", "A1", "资产/权益价值；各行明确是否已经扣除相关负债，避免重复扣减")
        terms = {case: [] for case in CASES}
        for i, item in enumerate(data["valuation"].get("assets", []), 4):
            graph.set("NAV", f"A{i}", item["name"])
            for col, case, key in (("B", "bear", "bear_value"), ("C", "base", "value"), ("D", "bull", "bull_value")):
                ref = raw(graph, data, "NAV", f"{col}{i}", item.get(key), f"valuation.assets.{i-4}.{key}")
                terms[case].append(ref)
        if not terms["base"]:
            raise ValueError("NAV需要资产价值项目")
        for col, label in (("B", "Bear资产价值"), ("C", "Base资产价值"), ("D", "Bull资产价值")):
            graph.set("NAV", f"{col}3", label)
        for i, case in enumerate(CASES, 4):
            graph.set("NAV", f"F{i}", f"{case} equity")
            case_values[case] = formula(graph, "NAV", f"G{i}", bridge(refs, "+".join(terms[case])))
    for i, case in enumerate(CASES, 22):
        graph.set("Valuation", f"A{i}", f"{case} equity value")
        if case_values:
            formula(graph, "Valuation", f"B{i}", case_values[case])
    graph.set("Valuation", "A26", "Selected equity value")
    graph.set("Valuation", "A27", "Selected per share value")
    if case_values:
        formula(graph, "Valuation", "B26", "CHOOSE(B2,B22,B23,B24)")
        formula(graph, "Valuation", "B27", f"B26/{refs['diluted_shares']}*{refs['scale']}")
    else:
        graph.set("Valuation", "B26", "不可估值：" + data["valuation"].get("reason", "关键资料不足"))
    core_sheets = {s for s, _ in graph.cells}
    for schedule in data.get("supplemental_schedules", []):
        sheet = schedule["name"]
        import re
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,30}", sheet) or sheet in core_sheets:
            raise ValueError("行业附表需唯一英文短名，不得覆盖核心表")
        core_sheets.add(sheet)
        for cell in schedule["cells"]:
            address = cell["address"]
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", address) or (sheet, address) in graph.cells:
                raise ValueError("行业附表地址无效或重复")
            if "formula" in cell:
                graph.set(sheet, address, formula=cell["formula"], note=cell.get("note", ""))
            elif isinstance(cell.get("value"), str):
                graph.set(sheet, address, cell["value"])
            else:
                raw(graph, data, sheet, address, cell.get("value"), cell.get("input_path", f"supplemental.{sheet}.{address}"))
    graph.set("Notes", "A1", "方法、来源、限制")
    notes = [f"币种 {data['currency']}；金额 {data['money_unit']}；股数 {data['shares_unit']}",
             "空白=未知；蓝字=输入，黑字=公式，绿字=跨表公式；合成算例不是事实。",
             "普通企业基线：年末折现、期初债务利息、无自动融资，其他资产负债冻结。",
             "SBC在FCFF保留为经济费用；未来同一授予不再增加估值股数；复杂稀释需另建模。",
             "三表平衡不证明预测经济合理；融资缺口、终值再投资、行业附表必须人工复核。",
             "NAV缺少Bear/Bull资产情景时保留未知，不把Base重复显示为三情景。",
             "公式Python计算结果保存在model-results.json；Excel缓存需实际引擎重算。"] + data.get("limitations", [])
    for i, note in enumerate(notes, 3):
        graph.set("Notes", f"A{i}", note)
    return graph


def calculate(graph, data):
    values, errors, missing, scope_limitations = {}, [], [], []
    for (sheet, address), cell in graph.cells.items():
        if cell.formula:
            try:
                values[f"{sheet}!{address}"] = graph.get(sheet, address)
            except MissingInput as exc:
                values[f"{sheet}!{address}"] = None
                missing.append(str(exc))
            except (ValueError, ZeroDivisionError, OverflowError) as exc:
                errors.append(f"{sheet}!{address}: {exc}")
    method = data["valuation"]["method"]
    checks = []
    def check(name, value, expected=0, tolerance=1e-6):
        checks.append(dict(name=name, value=value, expected=expected,
                           passed=value is not None and abs(value-expected) <= tolerance))
    for key, value in values.items():
        sheet, address = key.split("!")
        if sheet in {"Actuals", "Bear", "Base", "Bull"} and method == "fcff":
            row = int(''.join(x for x in address if x.isdigit()))
            if row in [ROWS[x] for x in ("balance_check", "cash_check", "re_check", "ppe_check")]:
                check(key, value)
            if row in [ROWS[x] for x in ("financing_gap",)]:
                check(key, value)
            if row in [ROWS[x] for x in ("cash", "debt", "ppe", "diluted_shares")]:
                if value is not None and value < 0:
                    errors.append(f"{key}: 负现金/债务/PP&E/股数，需处理融资或模型口径")
    v = data["valuation"]
    if method in {"fcff", "rim", "ddm"}:
        rate = v.get("wacc") if method == "fcff" else v.get("cost_equity")
        growth = v.get("terminal_growth")
        if rate is None or growth is None or rate <= growth or rate <= -1:
            errors.append("折现率须大于终值增长率且大于-1")
    if v.get("diluted_shares") is not None and v["diluted_shares"] <= 0:
        errors.append("估值股数必须大于零；未知可仅报告权益价值")
    if method == "fcff":
        center = values.get("Sensitivity!D6")
        base = values.get("Valuation!B23")
        if base is not None and v.get("diluted_shares"):
            check("sensitivity_center_vs_base", center, base/v["diluted_shares"]*SCALES[data["money_unit"]]/SCALES[data["shares_unit"]])
        # Baseline cannot safely value unmodeled future equity financing/dilution.
        for case in CASES:
            for period in data["scenarios"][case]:
                if period.get("equity_issued", 0) or period.get("shares_issued", 0) or period.get("shares_repur", 0):
                    errors.append("未来股权融资/股数变化需另建每股价值分配，基线静态股数估值不适用")
                    break
    if method == "rim" and all(v.get(x) is not None for x in ("terminal_growth", "terminal_roe", "terminal_payout")):
        check("rim_terminal_growth_supported_by_retention", v["terminal_growth"], v["terminal_roe"]*(1-v["terminal_payout"]), 1e-5)
    for item in data.get("industry_checks", []):
        key = f"{item['sheet']}!{item['address']}"
        try:
            value = graph.get(item["sheet"], item["address"])
        except (MissingInput, ValueError, ZeroDivisionError):
            value = None
        check(item["name"], value, item.get("expected", 0), item.get("tolerance", 1e-6))
    if data["industry_type"] in {"bank", "insurance", "asset", "early_growth"} and method != "unavailable":
        if not data.get("supplemental_schedules") or not data.get("industry_checks"):
            scope_limitations.append("该行业仍需有来源的经营/资本/融资附表及行业检查；基本估值构件不足以完成行业研究")
    untraced = [f"{s}!{a}" for (s,a), c in graph.cells.items() if c.note.startswith("UNTRACED") and c.value is not None]
    status = "limited" if method == "unavailable" or missing or untraced or errors or scope_limitations or any(not c["passed"] for c in checks) else "calculated"
    return dict(company=data["company"], as_of=data["as_of"], method=method, status=status,
                formula_values=values, checks=checks, errors=sorted(set(errors)), missing=sorted(set(missing)),
                untraced_inputs=untraced, scope_limitations=scope_limitations, recalculation="python_only; Excel engine not yet verified")


def export_workbook(graph, data, output):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.properties import CalcProperties
    book = Workbook()
    book.remove(book.active)
    sheets = list(dict.fromkeys(["Valuation"] + [s for s,_ in graph.cells]))
    for name in sheets:
        ws = book.create_sheet(name)
        ws.freeze_panes = "B5"
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 36 if name != "Notes" else 110
        for i in range(2, 16):
            ws.column_dimensions[column(i)].width = 19
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    for (sheet, address), item in graph.cells.items():
        cell = book[sheet][address]
        cell.value = item.formula if item.formula else ("=NA()" if item.value is None else item.value)
        if isinstance(item.value, str) and not item.formula:
            cell.data_type = "s"
        cell.font = Font(name="Calibri", size=11, color="008000" if item.formula and "!" in item.formula else ("000000" if item.formula else "0000CC"))
        cell.number_format = "0.0%;(0.0%);–" if item.percent else '#,##0.00;(#,##0.00);"–"'
        if item.note:
            cell.comment = Comment(("MISSING: " if item.value is None and not item.formula else "") + item.note, "Source / assumption")
        if cell.row == 1:
            cell.fill = PatternFill("solid", fgColor="183448")
            cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            book[sheet].row_dimensions[1].height = 45
        elif cell.column == 1:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if item.value is None and not item.formula and cell.column > 1:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    if "Assumptions" in book:
        validation = DataValidation(type="list", formula1='"1,2,3"')
        book["Assumptions"].add_data_validation(validation)
        validation.add(book["Assumptions"]["B2"])
    for ws in book:
        ws.merge_cells("A1:F1")
        if ws.title == "Notes":
            for row in range(3, ws.max_row + 1):
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
                ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[row].height = 40
    if "Sensitivity" in book:
        book["Sensitivity"]["D6"].fill = PatternFill("solid", fgColor="BDD7EE")
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


def build(data, output):
    graph = make_graph(data)
    result = calculate(graph, data)
    export_workbook(graph, data, output)
    result["input_sha256"] = hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    result["workbook_sha256"] = hashlib.sha256(Path(output).read_bytes()).hexdigest()
    target = Path(output).parent / "model-results.json"
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--summary", action="store_true", help="Print counts/examples; retain full model-results.json")
    args = parser.parse_args()
    try:
        result = build(json.loads(args.input.read_text()), args.output)
    except (ValueError, KeyError, ImportError) as exc:
        parser.exit(2, str(exc) + "\n")
    if args.summary:
        from research_context import summarize
        report = summarize(result, args.output.parent / "model-results.json",
                           ("errors", "missing", "untraced_inputs", "scope_limitations"))
    else:
        report = {k: result[k] for k in ("status", "errors", "missing", "untraced_inputs")}
    print(json.dumps(report, ensure_ascii=False))
    raise SystemExit(1 if result["errors"] else 0)
