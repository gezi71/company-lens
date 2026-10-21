#!/usr/bin/env python3
"""Recalculate in an isolated LibreOffice profile and compare cached formula values."""
import argparse
import hashlib
import json
import math
import shutil
import subprocess
import tempfile
from pathlib import Path


def check(path, recalculate=False):
    from openpyxl import load_workbook
    path = Path(path).resolve()
    result_path = path.parent / "model-results.json"
    result = json.loads(result_path.read_text())
    if hashlib.sha256(path.read_bytes()).hexdigest() != result.get("workbook_sha256"):
        raise ValueError("工作簿已改变；先按当前输入重建Python计算结果，再核对缓存")
    source = load_workbook(path, data_only=False)
    expected = result["formula_values"]
    report = dict(engine="not_run", errors=[], missing=[], compared=0)
    if not recalculate:
        report["engine"] = "not_requested"
        return report
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        report["engine"] = "unavailable"
        result["engine_check"] = report
        result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
        return report
    with tempfile.TemporaryDirectory(prefix="company-model-") as temp:
        temp = Path(temp)
        input_dir, output_dir = temp / "in", temp / "out"
        input_dir.mkdir()
        output_dir.mkdir()
        shutil.copy2(path, input_dir / path.name)
        profile = (temp / "profile").as_uri()
        command = [executable, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(input_dir / path.name)]
        process = subprocess.run(command, capture_output=True, text=True, timeout=60)
        rebuilt = output_dir / path.name
        if process.returncode or not rebuilt.exists():
            report["engine"] = "failed"
            report["errors"].append((process.stderr or process.stdout or "未生成重算文件")[:1000])
        else:
            cached = load_workbook(rebuilt, data_only=True)
            for ws in source:
                for row in ws:
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        key = f"{ws.title}!{cell.coordinate}"
                        value = cached[ws.title][cell.coordinate].value
                        target = expected.get(key)
                        intentional_missing = (cell.value == "=NA()" and cell.comment and cell.comment.text.startswith("MISSING:"))
                        if intentional_missing or (key in expected and target is None):
                            if value != "#N/A":
                                report["errors"].append(f"{key}: 未知值应传播#N/A，实际 {value!r}")
                            report["missing"].append(key)
                            continue
                        if key not in expected:
                            report["errors"].append(f"{key}: 缺少Python对照计算")
                        elif isinstance(target, (int, float)):
                            if not isinstance(value, (int, float)) or not math.isclose(value, target, rel_tol=1e-8, abs_tol=1e-6):
                                report["errors"].append(f"{key}: 引擎{value!r} / Python{target!r}")
                            else:
                                report["compared"] += 1
                        elif target != value:
                            report["errors"].append(f"{key}: 引擎{value!r} / Python{target!r}")
            report["engine"] = "verified_with_missing" if report["missing"] else "verified"
            if report["errors"]:
                report["engine"] = "mismatch"
            else:
                shutil.copy2(rebuilt, path)
                result["workbook_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    result["engine_check"] = report
    result["recalculation"] = report["engine"]
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--recalculate", action="store_true")
    parser.add_argument("--summary", action="store_true", help="Print counts/examples; full engine_check remains in model-results.json")
    args = parser.parse_args()
    try:
        report = check(args.workbook, args.recalculate)
    except (ValueError, OSError, subprocess.TimeoutExpired) as exc:
        parser.exit(2, str(exc) + "\n")
    if args.summary:
        from research_context import summarize
        output = summarize(report, args.workbook.parent / "model-results.json", ("errors", "missing"))
    else:
        output = report
    print(json.dumps(output, ensure_ascii=False))
    raise SystemExit(1 if report["errors"] else 0)
